"""Build docs/READOUT.pptx: a 5-slide leadership readout (objective -> assumptions -> findings ->
recommendation -> next steps). Numbers come from render_docs.values(), the same source as the memo.
Opens in PowerPoint, Keynote and Google Slides (no macros, standard fonts, plain shapes).
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from render_docs import values

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "READOUT.pptx"

INK, PAPER, ACCENT, MUTED, ORANGE = (RGBColor(0x0F, 0x11, 0x13), RGBColor(0xF4, 0xF1, 0xEA),
                                     RGBColor(0x19, 0x9E, 0x70), RGBColor(0x9A, 0xA0, 0xA6), RGBColor(0xD9, 0x59, 0x26))
W, H = Emu(12192000), Emu(6858000)   # 16:9
SERIF, SANS = "Georgia", "Arial"      # safe fonts that exist in Google Slides


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, s, size, color, font=SANS, bold=False):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.space_after = Pt(8)
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.name, r.font.bold = Pt(size), color, font, bold
    return box


def eyebrow(slide, s, color):
    text(slide, 600000, 450000, 11000000, 400000, s.upper(), 11, color)


def main():
    v = values()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # 1. Objective
    s = prs.slides.add_slide(blank)
    background(s, INK)
    eyebrow(s, "Ads analytics for a delivery marketplace · public data, not Swiggy data", MUTED)
    text(s, 600000, 1300000, 10500000, 2200000, "Three moves for the ads P&L", 54, PAPER, SERIF)
    text(s, 600000, 3500000, 10000000, 1500000,
         ["Objective: find where ad revenue can grow without hurting advertisers or users.",
          f"Data: {v['bids']} auction bids and {v['impressions']} impressions (iPinYou, {v['first_day']} to {v['last_day']}), "
          f"{v['campaigns']} Criteo campaigns, a randomised test on {v['ab_users']} users."], 18, MUTED)
    text(s, 600000, 6000000, 6000000, 400000, "Pranay Sarkar", 14, PAPER)

    # 2. Assumptions
    s = prs.slides.add_slide(blank)
    background(s, PAPER)
    eyebrow(s, "What I assumed", RGBColor(0x5B, 0x61, 0x68))
    text(s, 600000, 1000000, 11000000, 900000, "The data stands in for a marketplace's own ad slots", 34, INK, SERIF)
    text(s, 600000, 2300000, 10800000, 3800000, [
        "Auction logs behave like our slots: highest bid wins, winner pays the second price or the floor.",
        f"Criteo's {v['campaigns']} campaigns stand in for our advertisers; costs are compared, never read as currency.",
        "Floor model keeps bids fixed. The unknown, how many at-risk impressions still sell, is shown as a break-even, not guessed.",
        "The A/B decision rule was written before the readout.",
        "Fill rate and ROAS can't be computed from public data: win rate and cost per conversion are used instead.",
    ], 18, INK)

    # 3. Findings: four big numbers
    s = prs.slides.add_slide(blank)
    background(s, INK)
    eyebrow(s, "What I found", MUTED)
    text(s, 600000, 1000000, 11000000, 900000, "Thin auctions, an expensive advertiser group, ads that work", 32, PAPER, SERIF)
    cards = [
        (v["floor_binding"], "of floored auctions: the floor set the price", f"Exchange {v['no_floor_exchange']} sells {v['no_floor_share']} of impressions with no floor"),
        (v["x30_breakeven"], "must still sell for a 30-fen floor to pay off", f"{v['x30_at_risk']} of impressions at risk; range {v['x30_lower']} to {v['x30_upper']}"),
        (v["fix_cpa_ratio"], "cost per conversion: 'fix' vs 'protect' group", f"'Fix' = {v['fix_share']} of spend, week-4 spend {v['fix_retention']} of week 1"),
        (v["conv_rel"], "conversion lift from showing ads", f"95% CI {v['conv_rel_lo']} to {v['conv_rel_hi']}; decision {v['decision']}"),
    ]
    cw = 2650000
    for i, (big, label, sub) in enumerate(cards):
        x = 600000 + i * (cw + 120000)
        text(s, x, 2400000, cw, 1100000, big, 48, ACCENT if i != 2 else ORANGE, SERIF)
        text(s, x, 3500000, cw, 900000, label, 15, PAPER, bold=True)
        text(s, x, 4400000, cw, 1300000, sub, 13, MUTED)

    # 4. Recommendation
    s = prs.slides.add_slide(blank)
    background(s, PAPER)
    eyebrow(s, "Recommendation", RGBColor(0x5B, 0x61, 0x68))
    text(s, 600000, 1000000, 11000000, 900000, "Fix first, test the floor, grow the efficient", 34, INK, SERIF)
    recs = [
        ("1  Fix before we price", f"Targeting and bid help for the {v['fix_n']} 'fix' advertisers: {v['fix_share']} of spend, already shrinking."),
        ("2  Test, don't roll out, a minimum floor", f"A/B test on the exchange with no floors. Ship only if retention clears {v['x30_breakeven']}."),
        ("3  Pitch budget to the 'grow' group", f"{v['grow_n']} advertisers convert cheapest but spend {v['grow_share']} of the total."),
    ]
    for i, (head, body) in enumerate(recs):
        y = 2300000 + i * 1250000
        text(s, 600000, y, 10800000, 500000, head, 22, INK, SERIF, bold=False)
        text(s, 600000, y + 480000, 10800000, 700000, body, 16, RGBColor(0x5B, 0x61, 0x68))

    # 5. Next steps and limits
    s = prs.slides.add_slide(blank)
    background(s, INK)
    eyebrow(s, "Next steps and limits", MUTED)
    text(s, 600000, 1000000, 11000000, 900000, "What I would do with real marketplace data", 32, PAPER, SERIF)
    text(s, 600000, 2200000, 5300000, 3800000, [
        "Next",
        "Floor A/B test: measure the share that still sells.",
        "Single-campaign holdout to confirm the ad lift.",
        "Weekly KPI report from the same SQL and metric definitions.",
        "Add fill rate and ROAS once request logs and order values exist.",
    ], 16, PAPER)
    text(s, 6300000, 2200000, 5300000, 3800000, [
        "Limits",
        "2013 exchange prices, not rupees; Criteo costs are transformed.",
        "iPinYou totals after Oct 21 are lower bounds (test set withheld).",
        "A/B data pools several tests; effect on exposed users not used.",
        "Descriptive findings, except the randomised A/B test.",
    ], 16, MUTED)

    prs.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
