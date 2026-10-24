"""Build excel/ads_kpi_model.xlsx from the outputs/ CSVs.

Data sheets hold values; every summary sheet is live formulas (SUM, SUMIFS, COUNTIFS, XLOOKUP,
NORM.S.INV, NORM.S.DIST, IFERROR), so changing an input recalculates everything.
Only functions Google Sheets also supports are used: no macros, pivot objects or Power Query.
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs"
XLSX = ROOT / "excel" / "ads_kpi_model.xlsx"

INK, PAPER, ACCENT, INPUT = "0F1113", "F4F1EA", "2E9E83", "FFF4CC"
HEADER = dict(font=Font(bold=True, color=PAPER), fill=PatternFill("solid", fgColor=INK))
INPUT_STYLE = dict(fill=PatternFill("solid", fgColor=INPUT), font=Font(bold=True, color="7A5B00"))
TITLE = Font(bold=True, size=14)
NOTE = Font(italic=True, color="5B6168")
THIN = Border(bottom=Side(style="thin", color="D0CCC2"))

EXCHANGE_LABEL = {"ALL": "All", "1": "Exchange 1", "2": "Exchange 2", "3": "Exchange 3",
                  "4": "Exchange 4", "mobile/unknown": "Mobile/unknown"}


def read(name):
    with open(OUT / f"{name}.csv") as f:
        return list(csv.DictReader(f))


def num(v):
    """CSV text -> int/float where possible, '' -> None, else keep text."""
    if v in ("", None):
        return None
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return v


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HEADER["font"], HEADER["fill"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def data_sheet(wb, title, rows, formats=None, transform=None):
    """Write a CSV as a plain table. Returns the sheet."""
    ws = wb.create_sheet(title)
    cols = list(rows[0].keys())
    ws.append(cols)
    for r in rows:
        r = transform(dict(r)) if transform else r
        ws.append([num(r[c]) for c in cols])
    style_header(ws, 1, len(cols))
    ws.freeze_panes = "A2"
    for c, fmt in (formats or {}).items():
        idx = cols.index(c) + 1
        for row in range(2, len(rows) + 2):
            ws.cell(row=row, column=idx).number_format = fmt
    widths(ws, [max(12, min(28, len(c) + 4)) for c in cols])
    return ws


def label_rows(ws, start_row, items, col=1):
    """items = [(label, formula, number_format, note)] written down two columns."""
    for i, (label, formula, fmt, note) in enumerate(items):
        r = start_row + i
        ws.cell(row=r, column=col, value=label).border = THIN
        cell = ws.cell(row=r, column=col + 1, value=formula)
        cell.number_format, cell.border = fmt, THIN
        if note:
            ws.cell(row=r, column=col + 2, value=note).font = NOTE


def readme(wb, numbers):
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Ads KPI model: ads analytics for a delivery marketplace, built on public ad data", TITLE),
        ("Not Swiggy data. Sources: iPinYou RTB logs (season 3, Oct 2013), Criteo Attribution, Criteo Uplift v2.1.", None),
        ("Money: iPinYou in CNY (2013). Criteo cost is a transformed unit; compare campaigns, don't read it as currency.", None),
        ("Oct 21-28 iPinYou totals are lower bounds (iPinYou withheld a test set); ratios are fair.", None),
        ("", None),
        ("Sheets", Font(bold=True)),
        ("KPI_Summary: totals and ratios from the daily table (SUM), plus per-advertiser KPIs (XLOOKUP).", None),
        ("Segments: pivot-style summary of 675 campaigns (COUNTIFS / SUMIFS).", None),
        ("Campaign_Lookup: type a campaign id, get its numbers and segment (XLOOKUP).", None),
        ("Floor_Model: pick an exchange, a minimum floor and a retention rate; revenue and break-even update.", None),
        ("AB_Test: A/B readout recomputed with formulas (NORM.S.INV, NORM.S.DIST).", None),
        ("Daily, Advertisers, Campaigns, FloorData: values exported from the SQL layer.", None),
        ("", None),
        ("Yellow cells are inputs. Everything else is a formula or exported data.", Font(bold=True, color="7A5B00")),
        ("Google Sheets: Drive > New > File upload > open with Google Sheets. All functions used are supported.", None),
        (f"Built from outputs/numbers.json; iPinYou window {numbers['ipinyou_kpi']['first_day']} to {numbers['ipinyou_kpi']['last_day']}.", NOTE),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
    ws.column_dimensions["A"].width = 120


def kpi_summary(wb, n_daily, adv_ids):
    ws = wb.create_sheet("KPI_Summary")
    ws["A1"], ws["A1"].font = "Ads P&L KPIs (iPinYou, all days)", TITLE
    last = n_daily + 1
    items = [
        ("Bids", f"=SUM(Daily!B2:B{last})", "#,##0", ""),
        ("Bids won", f"=SUM(Daily!C2:C{last})", "#,##0", ""),
        ("Win rate", "=B4/B3", "0.0%", "stand-in for fill rate (bidder's view)"),
        ("Impressions", f"=SUM(Daily!E2:E{last})", "#,##0", "includes impressions with no matching bid"),
        ("Spend / revenue (CNY)", f"=SUM(Daily!F2:F{last})", "#,##0.00", "advertiser spend = marketplace ad revenue"),
        ("CPM (CNY per 1,000 imps)", "=B7*1000/B6", "0.000", "ratio of totals, not average of daily CPMs"),
        ("Clicks", f"=SUM(Daily!H2:H{last})", "#,##0", ""),
        ("CTR", "=B9/B6", "0.000%", ""),
        ("CPC (CNY)", "=B7/B9", "0.000", ""),
        ("Conversions", f"=SUM(Daily!K2:K{last})", "#,##0", "landing-page visits, 2 advertisers only"),
        ("Best day by spend", f"=INDEX(Daily!A2:A{last},MATCH(MAX(Daily!F2:F{last}),Daily!F2:F{last},0))", "@", "INDEX/MATCH"),
    ]
    label_rows(ws, 3, items)

    ws["A16"], ws["A16"].font = "Per advertiser (XLOOKUP into Advertisers)", Font(bold=True)
    head = ["Advertiser id", "Category", "Spend (CNY)", "Spend share", "CPM (CNY)", "CTR", "Win rate"]
    for c, h in enumerate(head, start=1):
        ws.cell(row=17, column=c, value=h)
    style_header(ws, 17, len(head))
    cols = {"Category": "B", "Spend (CNY)": "G", "CPM (CNY)": "I", "CTR": "K", "Win rate": "E"}
    for i, adv in enumerate(adv_ids):
        r = 18 + i
        ws.cell(row=r, column=1, value=adv)
        for c, h in enumerate(head[1:], start=2):
            if h == "Spend share":
                f = f"=C{r}/SUM($C$18:$C${17 + len(adv_ids)})"
            else:
                f = f"=_xlfn.XLOOKUP($A{r},Advertisers!$A:$A,Advertisers!${cols[h]}:${cols[h]},\"not found\")"
            cell = ws.cell(row=r, column=c, value=f)
            cell.number_format = {"Spend (CNY)": "#,##0.00", "Spend share": "0.0%", "CPM (CNY)": "0.000",
                                  "CTR": "0.000%", "Win rate": "0.0%"}.get(h, "General")
    widths(ws, [28, 18, 16, 14, 14, 12, 12])


def segments(wb, n_camp):
    ws = wb.create_sheet("Segments")
    ws["A1"], ws["A1"].font = "Advertiser segments: pivot-style summary of 675 Criteo campaigns", TITLE
    ws["A2"], ws["A2"].font = "Split at median spend and median CPA. Built with COUNTIFS/SUMIFS over the Campaigns sheet.", NOTE
    head = ["Segment", "Campaigns", "Spend", "Spend share", "Conversions", "CTR", "CPA", "Spend retention (wk4/wk1)",
            "Campaigns whose spend halved", "Action"]
    for c, h in enumerate(head, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(head))
    rng = lambda col: f"Campaigns!${col}$2:${col}${n_camp + 1}"
    seg, imp, clk, conv, spend = rng("O"), rng("C"), rng("D"), rng("E"), rng("F")
    actions = {"PROTECT": "big and efficient: keep, upsell", "FIX": "big but expensive conversions: optimise before they cut spend",
               "GROW": "efficient but small: pitch more budget", "WATCH": "small and expensive: low priority"}
    for i, s in enumerate(["PROTECT", "FIX", "GROW", "WATCH"]):
        r = 5 + i
        ws.cell(row=r, column=1, value=s).font = Font(bold=True)
        ws.cell(row=r, column=2, value=f"=COUNTIFS({seg},A{r})")
        ws.cell(row=r, column=3, value=f"=SUMIFS({spend},{seg},A{r})").number_format = "#,##0.00"
        ws.cell(row=r, column=4, value=f"=C{r}/$C$9").number_format = "0.0%"
        ws.cell(row=r, column=5, value=f"=SUMIFS({conv},{seg},A{r})").number_format = "#,##0"
        ws.cell(row=r, column=6, value=f"=SUMIFS({clk},{seg},A{r})/SUMIFS({imp},{seg},A{r})").number_format = "0.0%"
        ws.cell(row=r, column=7, value=f"=C{r}/E{r}").number_format = "0.0000"
        ws.cell(row=r, column=8, value=f"=SUMIFS(Campaigns!$N$2:$N${n_camp + 1},{seg},A{r})/SUMIFS(Campaigns!$M$2:$M${n_camp + 1},{seg},A{r})").number_format = "0.00"
        ws.cell(row=r, column=9, value=f"=COUNTIFS({seg},A{r},Campaigns!$L$2:$L${n_camp + 1},\"<0.5\")")
        ws.cell(row=r, column=10, value=actions[s]).font = NOTE
    ws.cell(row=9, column=1, value="Total").font = Font(bold=True)
    for c, f, fmt in [(2, "=SUM(B5:B8)", "0"), (3, "=SUM(C5:C8)", "#,##0.00"), (4, "=SUM(D5:D8)", "0.0%"),
                      (5, "=SUM(E5:E8)", "#,##0"), (6, f"=SUM({clk})/SUM({imp})", "0.0%"), (7, "=C9/E9", "0.0000"),
                      (8, f"=SUM(Campaigns!$N$2:$N${n_camp + 1})/SUM(Campaigns!$M$2:$M${n_camp + 1})", "0.00"),
                      (9, "=SUM(I5:I8)", "0")]:
        cell = ws.cell(row=9, column=c, value=f)
        cell.number_format, cell.font = fmt, Font(bold=True)
    ws["A11"] = "FIX vs PROTECT: cost per conversion is"
    ws["B11"] = "=G6/G5"
    ws["B11"].number_format = '0.0"x"'
    ws["C11"] = "higher for FIX, while FIX holds this share of spend:"
    ws["D11"] = "=D6"
    ws["D11"].number_format = "0.0%"
    ws["A12"] = "Top 20% of campaigns hold this share of spend:"
    ws["B12"] = f"=SUMIFS({spend},Campaigns!$B$2:$B${n_camp + 1},\"<=\"&ROUND(0.2*B9,0))/C9"
    ws["B12"].number_format = "0.0%"
    widths(ws, [44, 12, 12, 12, 13, 9, 10, 16, 16, 58])


def campaign_lookup(wb, n_camp, default_id):
    ws = wb.create_sheet("Campaign_Lookup")
    ws["A1"], ws["A1"].font = "Look up any campaign", TITLE
    ws["A3"] = "Campaign id (type one)"
    ws["B3"] = default_id
    for k, v in INPUT_STYLE.items():
        setattr(ws["B3"], k, v)
    rng = lambda col: f"Campaigns!${col}$2:${col}${n_camp + 1}"
    items = [
        ("Spend rank (1 = biggest)", "B", "0"), ("Spend", "F", "#,##0.0000"), ("Share of all spend", "G", "0.00%"),
        ("Conversions", "E", "#,##0"), ("CTR", "I", "0.0%"), ("CVR (conversions / clicks)", "J", "0.0%"),
        ("CPA (cost per conversion)", "K", "0.000000"), ("Spend retention wk4 / wk1", "L", "0.00"), ("Segment", "O", "@"),
    ]
    for i, (label, col, fmt) in enumerate(items):
        r = 5 + i
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=f"=_xlfn.XLOOKUP($B$3,{rng('A')},{rng(col)},\"no such campaign\")")
        c.number_format = fmt
    ws["A15"] = "Median CPA of all campaigns"
    ws["B15"] = f"=MEDIAN({rng('K')})"
    ws["B15"].number_format = "0.000000"
    ws["A16"] = "This campaign's CPA vs median"
    ws["B16"] = '=IFERROR(B11/B15,"-")'
    ws["B16"].number_format = '0.00"x"'
    widths(ws, [34, 22])


def floor_model(wb, n_floor):
    ws = wb.create_sheet("Floor_Model")
    ws["A1"], ws["A1"].font = "Minimum floor price: what would it earn?", TITLE
    ws["A2"], ws["A2"].font = ("New floor = max(current floor, X). Impressions that paid less than X are at risk: "
                               "a share r still sells at X, the rest is lost (ADR 0007). r is unknown; an A/B test measures it."), NOTE
    inputs = [("Exchange", "All"), ("Minimum floor X (fen per 1,000 imps)", 30), ("Retention r (share of at-risk imps that still sell)", 0.5)]
    for i, (label, val) in enumerate(inputs):
        ws.cell(row=4 + i, column=1, value=label)
        c = ws.cell(row=4 + i, column=2, value=val)
        for k, v in INPUT_STYLE.items():
            setattr(c, k, v)
    ws["B6"].number_format = "0%"
    dv_ex = DataValidation(type="list", formula1='"' + ",".join(EXCHANGE_LABEL.values()) + '"', allow_blank=False)
    dv_x = DataValidation(type="list", formula1='"' + ",".join(str(x) for x in range(0, 160, 10)) + '"')
    dv_r = DataValidation(type="decimal", operator="between", formula1="0", formula2="1")
    for dv, ref in [(dv_ex, "B4"), (dv_x, "B5"), (dv_r, "B6")]:
        ws.add_data_validation(dv)
        dv.add(ref)

    f = lambda col: f"SUMIFS(FloorData!${col}$2:${col}${n_floor + 1},FloorData!$A$2:$A${n_floor + 1},$B$4,FloorData!$B$2:$B${n_floor + 1},$B$5)"
    items = [
        ("Revenue today (CNY)", f"={f('G')}", "#,##0.00", ""),
        ("Impressions at risk (paid < X)", f"={f('D')}", "#,##0", ""),
        ("Share of impressions at risk", f"=B10/{f('C')}", "0.0%", ""),
        ("Revenue if every at-risk imp is lost (r = 0)", f"={f('H')}", "#,##0.00", "lower bound"),
        ("Revenue at your r", f"=B12+$B$6*$B$5*{f('E')}/100000", "#,##0.00", "safe revenue + r × X × sellable at-risk imps"),
        ("Change vs today", "=B13/B9-1", "+0.0%;-0.0%;0.0%", ""),
        ("Break-even r", f"=IFERROR((B9-B12)/($B$5*{f('E')}/100000),\"n/a\")", "0.0%", "below this r, the floor loses money"),
        ("Verdict at your r", '=IF(B5=0,"no change",IF(B13>=B9,"raises revenue","loses revenue"))', "@", ""),
    ]
    label_rows(ws, 9, items)

    ws["A19"], ws["A19"].font = "Sensitivity: revenue change vs today, by minimum floor (rows) and retention r (columns)", Font(bold=True)
    rs = [0, 0.25, 0.5, 0.75, 1]
    ws.cell(row=20, column=1, value="X \\ r")
    for j, r in enumerate(rs):
        ws.cell(row=20, column=2 + j, value=r).number_format = "0%"
    style_header(ws, 20, 1 + len(rs))
    g = lambda col, xref: (f"SUMIFS(FloorData!${col}$2:${col}${n_floor + 1},FloorData!$A$2:$A${n_floor + 1},$B$4,"
                           f"FloorData!$B$2:$B${n_floor + 1},{xref})")
    for i, x in enumerate(range(0, 160, 10)):
        row = 21 + i
        ws.cell(row=row, column=1, value=x)
        for j in range(len(rs)):
            col = get_column_letter(2 + j)
            cell = ws.cell(row=row, column=2 + j,
                           value=f"=({g('H', f'$A{row}')}+{col}$20*$A{row}*{g('E', f'$A{row}')}/100000)/{g('G', f'$A{row}')}-1")
            cell.number_format = "+0.0%;-0.0%;0.0%"
    widths(ws, [52, 16, 12, 12, 12, 12])


def ab_test(wb, ab):
    ws = wb.create_sheet("AB_Test")
    ws["A1"], ws["A1"].font = "A/B readout: does showing ads raise conversion? (Criteo Uplift, 13.98M users)", TITLE
    ws["A2"], ws["A2"].font = "Counts are inputs; every statistic below is a formula. Design fixed before the readout (docs/EXPERIMENT.md).", NOTE
    t, c = ab["counts"]["T"], ab["counts"]["C"]
    ws["A4"], ws["B4"], ws["C4"] = "", "Treatment", "Control"
    style_header(ws, 4, 3)
    rows = [("Users", t["users"], c["users"]), ("Visits", t["visits"], c["visits"]),
            ("Conversions", t["conversions"], c["conversions"]), ("Exposed to an ad", t["exposed"], c["exposed"])]
    for i, (label, tv, cv) in enumerate(rows):
        ws.cell(row=5 + i, column=1, value=label)
        for col, v in ((2, tv), (3, cv)):
            cell = ws.cell(row=5 + i, column=col, value=v)
            cell.number_format = "#,##0"
            for k, s in INPUT_STYLE.items():
                setattr(cell, k, s)
    ws["A10"], ws["B10"] = "Alpha (two-sided)", 0.05
    ws["A11"], ws["B11"] = "Intended treatment share", 0.85
    for ref in ("B10", "B11"):
        for k, s in INPUT_STYLE.items():
            setattr(ws[ref], k, s)
    items = [
        ("Observed treatment share", "=B5/(B5+C5)", "0.0000%", "SRM check: should match the intended share"),
        ("SRM chi-square", "=(B5-(B5+C5)*B11)^2/((B5+C5)*B11)+(C5-(B5+C5)*(1-B11))^2/((B5+C5)*(1-B11))", "0.000000", "compare with 10.83 (p = 0.001, 1 df)"),
        ("Conversion rate, treatment", "=B7/B5", "0.0000%", ""),
        ("Conversion rate, control", "=C7/C5", "0.0000%", ""),
        ("Absolute lift (percentage points)", "=B15-B16", "0.0000%", ""),
        ("Relative lift", "=B15/B16-1", "+0.0%", ""),
        ("z critical", "=_xlfn.NORM.S.INV(1-B10/2)", "0.0000", "1.96 for 95%"),
        ("Standard error (unpooled)", "=SQRT(B15*(1-B15)/B5+B16*(1-B16)/C5)", "0.000000", ""),
        ("95% CI low", "=B17-B19*B20", "0.0000%", ""),
        ("95% CI high", "=B17+B19*B20", "0.0000%", ""),
        ("Pooled rate", "=(B7+C7)/(B5+C5)", "0.0000%", "used only for the test, assuming no difference"),
        ("z statistic", "=B17/SQRT(B23*(1-B23)*(1/B5+1/C5))", "0.00", ""),
        ("p-value", "=2*(1-_xlfn.NORM.S.DIST(ABS(B24),TRUE))", "0.00E+00", "shows 0 when smaller than spreadsheet precision"),
        ("Visit rate, treatment", "=B6/B5", "0.00%", ""),
        ("Visit rate, control", "=C6/C5", "0.00%", ""),
        ("Visit lift CI low (guardrail)", "=(B26-B27)-B19*SQRT(B26*(1-B26)/B5+B27*(1-B27)/C5)", "0.0000%", "must be >= 0"),
        ("MDE, conversion (80% power)", "=(B19+_xlfn.NORM.S.INV(0.8))*SQRT(B16*(1-B16)*(1/B5+1/C5))", "0.0000%", "smallest lift this test could reliably detect"),
        ("Exposure rate in treatment", "=B8/B5", "0.00%", ""),
        ("Lift for users who saw an ad (CACE)", "=B17/B30", "0.00%", "ITT / exposure rate"),
        ("Decision (pre-registered rule)", '=IF(B14>=10.83,"untrustworthy (SRM)",IF(AND(B21>0,B28>=0),"SHIP",IF(OR(B22<0,B33<0),"DON\'T SHIP","INCONCLUSIVE")))', "@", "rule in docs/EXPERIMENT.md section 7"),
        ("Visit lift CI high (guardrail)", "=(B26-B27)+B19*SQRT(B26*(1-B26)/B5+B27*(1-B27)/C5)", "0.0000%", ""),
    ]
    label_rows(ws, 13, items)
    widths(ws, [40, 18, 18, 52])


def main():
    numbers = json.loads((OUT / "numbers.json").read_text())
    daily, adv, camp, floor = read("kpi_daily"), read("kpi_advertiser"), read("campaign_segments"), read("floor_sizing")
    wb = Workbook()
    readme(wb, numbers)
    kpi_summary(wb, len(daily), [int(a["advertiser_id"]) for a in adv])
    segments(wb, len(camp))
    campaign_lookup(wb, len(camp), int(camp[0]["campaign"]))
    floor_model(wb, len(floor))
    ab_test(wb, numbers["ab"])
    data_sheet(wb, "Daily", daily, {"win_rate": "0.0%", "ctr": "0.000%", "spend_cny": "#,##0.00"})
    data_sheet(wb, "Advertisers", adv, {"win_rate": "0.0%", "spend_share": "0.0%", "ctr": "0.000%"})
    data_sheet(wb, "Campaigns", camp, {"spend_share": "0.00%", "cumulative_spend_share": "0.0%"})
    data_sheet(wb, "FloorData", floor, transform=lambda r: {**r, "exchange": EXCHANGE_LABEL[r["exchange"]]})
    XLSX.parent.mkdir(exist_ok=True)
    wb.save(XLSX)
    print(f"wrote {XLSX.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
